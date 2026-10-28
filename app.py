import streamlit as st
import asyncio
from playwright.async_api import async_playwright
from docx import Document
from urllib.parse import urlparse, parse_qs
import io
import os

st.set_page_config(page_title="Auditor Privacy & Consent Mode v2", page_icon="🛡️", layout="wide")

@st.cache_resource
def ensure_playwright():
    os.system("playwright install chromium")

ensure_playwright()

st.title("🛡️ Privacy & Consent Mode v2 Auditor")
st.caption("Auditoría técnica automatizada: Matriz completa de consentimientos (5 fases)")

# Persistencia de memoria para evitar reseteos al interactuar con filtros
if "audit_data" not in st.session_state:
    st.session_state.audit_data = None
if "domain" not in st.session_state:
    st.session_state.domain = ""

target_url = st.text_input("Introduce la URL a auditar:", placeholder="https://es.bolognarthotels.it/")

CMP_CONFIG = {
    "CookieScript": {
        "accept": "#cookiescript_accept",
        "reject": "#cookiescript_reject",
        "manage": "#cookiescript_manage",
        "cat_perf": "#cookiescript_category_performance",
        "cat_target": "#cookiescript_category_targeting",
        "save": "#cookiescript_save"
    },
    "OneTrust": {
        "accept": "#onetrust-accept-btn-handler",
        "reject": "#onetrust-reject-all-handler",
        "manage": "#onetrust-pc-btn-handler",
        "cat_perf": "input[id*='ot-group-id-C0002']",
        "cat_target": "input[id*='ot-group-id-C0004']",
        "save": ".save-preference-btn-handler"
    },
    "Cookiebot": {
        "accept": "#CybotCookiebotDialogBodyLevelButtonLevelOptinAllowAll",
        "reject": "#CybotCookiebotDialogBodyButtonDecline",
        "manage": "#CybotCookiebotDialogBodyLevelButtonCustomize",
        "cat_perf": "#CybotCookiebotDialogBodyLevelButtonStatistics",
        "cat_target": "#CybotCookiebotDialogBodyLevelButtonMarketing",
        "save": "#CybotCookiebotDialogBodyLevelButtonLevelOptinAllowallSelection"
    }
}

def extract_consent_params(url):
    parsed = parse_qs(urlparse(url).query)
    return {
        'gcs': parsed.get('gcs', ['No detectado'])[0],
        'gcd': parsed.get('gcd', ['No detectado'])[0],
        'npa': parsed.get('npa', ['No detectado'])[0],
        'dma': parsed.get('dma', ['No detectado'])[0],
        'pscdl': parsed.get('pscdl', ['No detectado'])[0]
    }

async def interactuar_cmp(page, action):
    # 1. Acciones globales directas
    if action in ["accept", "reject"]:
        for cmp_name, cfg in CMP_CONFIG.items():
            sel = cfg.get(action)
            try:
                elem = page.locator(sel).first
                if await elem.is_visible(timeout=800):
                    await elem.click()
                    return f"{cmp_name} ({action})"
            except Exception:
                pass

        patterns = {
            "accept": ["Aceptar todas", "Aceptar todo", "Accept all", "Permitir todas", "Aceptar", "Accept"],
            "reject": ["Rechazar todas", "Rechazar todo", "Reject all", "Denegar", "Rechazar", "Reject"]
        }
        for txt in patterns.get(action, []):
            try:
                btn = page.locator(f"button:has-text('{txt}'), a:has-text('{txt}')").first
                if await btn.is_visible(timeout=800):
                    await btn.click()
                    return f"Heurística ('{txt}')"
            except Exception:
                pass

    # 2. Configuración granular en segunda capa
    elif action in ["only_analytics", "only_marketing"]:
        for cmp_name, cfg in CMP_CONFIG.items():
            manage_sel = cfg.get("manage")
            try:
                btn_manage = page.locator(manage_sel).first
                if await btn_manage.is_visible(timeout=1000):
                    await btn_manage.click()
                    await asyncio.sleep(1.5)
                    
                    perf_sel = cfg.get("cat_perf")
                    target_sel = cfg.get("cat_target")
                    save_sel = cfg.get("save")
                    
                    if action == "only_analytics":
                        if perf_sel and await page.locator(perf_sel).is_visible():
                            await page.locator(perf_sel).check()
                        if target_sel and await page.locator(target_sel).is_visible():
                            await page.locator(target_sel).uncheck()
                    elif action == "only_marketing":
                        if perf_sel and await page.locator(perf_sel).is_visible():
                            await page.locator(perf_sel).uncheck()
                        if target_sel and await page.locator(target_sel).is_visible():
                            await page.locator(target_sel).check()
                    
                    if save_sel:
                        btn_save = page.locator(save_sel).first
                        if await btn_save.is_visible(timeout=1200):
                            await btn_save.click()
                            return f"{cmp_name} ({action})"
            except Exception:
                pass

    return "No interactuado"

async def ejecutar_fase(browser, url, action=None):
    context = await browser.new_context()
    page = await context.new_page()
    pings = []
    
    page.on("request", lambda r: pings.append(r.url) if ("collect" in r.url and "google" in r.url) else None)
    
    try:
        await page.goto(url, timeout=20000, wait_until="domcontentloaded")
    except Exception:
        pass
    
    await asyncio.sleep(2)
    
    action_info = "Sin interacción (Default)"
    if action:
        action_info = await interactuar_cmp(page, action)
        await asyncio.sleep(2.5)
        
    cookies = await context.cookies()
    await context.close()
    return cookies, pings, action_info

def generar_documento_auditoria(domain, audit_data):
    doc = Document()
    doc.add_heading('AUDITORÍA PRIVACY', 0)
    doc.add_paragraph('Octubre 2026')
    doc.add_heading('Introducción.', 1)
    doc.add_paragraph('Este documento pretende revisar todas las cuestiones fundamentales referidas al cumplimiento de las premisas legales establecidas por el RGPD y las directrices de la AEPD...')
    
    doc.add_paragraph(f'\n{domain}')
    doc.add_heading('1.- HOME', 1)
    doc.add_heading('1.2 Revisión del consentimiento e inventario de cookies', 2)
    
    secciones = [
        ("1.2.1 Aterrizaje previo (Valores Default)", 'default'),
        ("1.2.2 Tras Rechazar Todas las cookies", 'reject'),
        ("1.2.3 Tras Aceptar Todas las cookies", 'accept'),
        ("1.2.4 Tras Aceptar Solo cookies Analíticas", 'analytics'),
        ("1.2.5 Tras Aceptar Solo cookies de Marketing / Publicitarias", 'marketing')
    ]
    for tit, k in secciones:
        doc.add_heading(tit, 3)
        cookies = audit_data[k]['cookies']
        if cookies:
            doc.add_paragraph(f"Total cookies detectadas: {len(cookies)}")
            for c in cookies:
                doc.add_paragraph(f"- {c['name']} (Dominio: {c['domain']})", style='List Bullet')
        else:
            doc.add_paragraph("No se detectaron cookies registradas.")
            
    doc.add_heading('1.3 Revisión de Consent Mode', 2)
    
    def add_cm_section(subtit, key, expected):
        doc.add_heading(subtit, 3)
        pings = audit_data[key]['pings']
        if pings:
            p = extract_consent_params(pings[-1] if key != 'default' else pings[0])
            doc.add_paragraph(f"- Señal GCS: {p['gcs']}")
            doc.add_paragraph(f"- Cadena GCD: {p['gcd']}")
            doc.add_paragraph(f"- Señales DMA y NPA: npa={p['npa']}, dma={p['dma']}")
            if expected in p['gcs']:
                doc.add_paragraph(f"Diagnóstico: Señal esperada ({expected}) transmitida correctamente.")
        else:
            doc.add_paragraph("Sin llamadas collect detectadas.")

    add_cm_section('1.3.1 Revisión previa al consentimiento (valores default)', 'default', 'G100')
    add_cm_section('1.3.2 Revisión tras rechazar todas (update)', 'reject', 'G100')
    add_cm_section('1.3.3 Revisión tras aceptar todas (valores update)', 'accept', 'G111')
    add_cm_section('1.3.4 Revisión tras aceptar publicitarias (valores update)', 'marketing', 'G110')
    add_cm_section('1.3.5 Revisión tras aceptar analíticas (valores update)', 'analytics', 'G101')

    doc.add_heading('2.- MOTOR DE RESERVA', 1)
    doc.add_paragraph('En este apartado se trata de observar si el consentimiento previamente dado se mantiene en el motor de reserva.')
    
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io

# Lanzador con barra de progreso
if st.button("🚀 Iniciar Auditoría Completa de Consentimiento", type="primary", use_container_width=True):
    if not target_url.startswith("http"):
        st.error("Introduce una URL completa con https:// o http://")
    else:
        st.session_state.domain = urlparse(target_url).netloc
        progreso = st.progress(0, text="Iniciando Chromium en la nube...")
        
        async def auditar_paso_a_paso():
            async with async_playwright() as p:
                browser = await p.chromium.launch(
                    headless=True,
                    args=["--no-sandbox", "--disable-setuid-sandbox", "--disable-dev-shm-usage"]
                )
                
                progreso.progress(10, text="[1/5] Fase 1: Aterrizaje Default...")
                c_def, p_def, a_def = await ejecutar_fase(browser, target_url, action=None)
                
                progreso.progress(30, text="[2/5] Fase 2: Rechazar Todas...")
                c_rej, p_rej, a_rej = await ejecutar_fase(browser, target_url, action="reject")
                
                progreso.progress(50, text="[3/5] Fase 3: Aceptar Todas...")
                c_acc, p_acc, a_acc = await ejecutar_fase(browser, target_url, action="accept")
                
                progreso.progress(70, text="[4/5] Fase 4: Solo Analítica...")
                c_ana, p_ana, a_ana = await ejecutar_fase(browser, target_url, action="only_analytics")
                
                progreso.progress(90, text="[5/5] Fase 5: Solo Marketing...")
                c_mkt, p_mkt, a_mkt = await ejecutar_fase(browser, target_url, action="only_marketing")
                
                await browser.close()
                progreso.progress(100, text="¡Completado con éxito!")
                
                return {
                    "default": {"cookies": c_def, "pings": p_def, "action": a_def},
                    "reject": {"cookies": c_rej, "pings": p_rej, "action": a_rej},
                    "accept": {"cookies": c_acc, "pings": p_acc, "action": a_acc},
                    "analytics": {"cookies": c_ana, "pings": p_ana, "action": a_ana},
                    "marketing": {"cookies": c_mkt, "pings": p_mkt, "action": a_mkt}
                }

        try:
            st.session_state.audit_data = asyncio.run(auditar_paso_a_paso())
            progreso.empty()
        except Exception as e:
            st.error(f"Error durante el escaneo: {e}")

# Renderizado condicional mediante session_state (persiste entre clics y filtros)
if st.session_state.audit_data:
    data = st.session_state.audit_data
    domain = st.session_state.domain
    
    st.success(f"¡Auditoría finalizada para {domain}!")
    
    t_cm, t_ck, t_doc = st.tabs(["📊 Google Consent Mode v2", "🍪 Cookies por Sección", "📄 Generar Informe Word"])
    
    with t_cm:
        st.markdown("### Señales recolectadas en las 5 fases")
        cols = st.columns(5)
        fases = [
            ("1. Default", 'default', "G100"),
            ("2. Rechazar", 'reject', "G100"),
            ("3. Aceptar Todo", 'accept', "G111"),
            ("4. Solo Analítica", 'analytics', "G101"),
            ("5. Solo Marketing", 'marketing', "G110")
        ]
        for col, (label, key, expected_gcs) in zip(cols, fases):
            with col:
                st.markdown(f"**{label}**")
                pings = data[key]['pings']
                if pings:
                    p = extract_consent_params(pings[-1] if key != 'default' else pings[0])
                    st.metric("GCS", p['gcs'])
                    if expected_gcs in p['gcs']:
                        st.caption(f"✅ Esperado: {expected_gcs}")
                    else:
                        st.caption(f"⚠️ Esperado: {expected_gcs}")
                    st.code(f"GCD: {p['gcd']}\nNPA: {p['npa']}\nDMA: {p['dma']}", language="text")
                else:
                    st.warning("Sin pings")
                st.caption(f"Acción: {data[key]['action']}")
    
    with t_ck:
        st.markdown("### Inventario de cookies por cada casuística")
        sec_names = {
            "default": "1. Aterrizaje Inicial (Default)",
            "reject": "2. Tras Rechazar Todas",
            "accept": "3. Tras Aceptar Todas",
            "analytics": "4. Tras Aceptar Solo Analítica",
            "marketing": "5. Tras Aceptar Solo Marketing"
        }
        sel_tab = st.selectbox(
            "Selecciona la casuística a inspeccionar:", 
            list(sec_names.keys()), 
            format_func=lambda x: sec_names[x],
            key="selected_cookie_category"
        )
        c_list = data[sel_tab]['cookies']
        st.write(f"Total cookies presentes: **{len(c_list)}**")
        if c_list:
            st.dataframe([{"Nombre": c['name'], "Dominio": c['domain'], "Expiración (Unix)": c['expires'], "Segura": c['secure']} for c in c_list], use_container_width=True)
        else:
            st.info("No hay cookies guardadas en esta fase.")

    with t_doc:
        st.markdown("### Descarga del documento oficial")
        doc_bytes = generar_documento_auditoria(domain, data)
        nombre_archivo = f"{domain} - Auditoría cookies y CMP.docx"
        st.download_button(
            label=f"📥 Descargar: {nombre_archivo}",
            data=doc_bytes,
            file_name=nombre_archivo,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary",
            key="download_doc_btn"
        )
